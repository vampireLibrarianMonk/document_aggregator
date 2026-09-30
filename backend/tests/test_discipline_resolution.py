"""Alpha-loop tests for the no-default discipline: effective-style resolution,
declared/learned/undefined precedence, and within-document formatting drift.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from app.convert.base import _collect_evidence  # noqa: E402
from app.convert.docx_in import extract_docx  # noqa: E402
from app.discipline import learn_discipline  # noqa: E402
from app.discipline.inspect import inspect as inspect_discipline  # noqa: E402
from app.discipline.spec import BuildDiscipline, TextFormat  # noqa: E402


def _doc(build) -> dict:
    import docx

    d = docx.Document()
    build(d)
    buf = io.BytesIO()
    d.save(buf)
    raw, _ = extract_docx(buf.getvalue())
    return _collect_evidence(raw)


# ---- effective-style resolution ----

def test_style_font_resolved_through_chain():
    """A heading styled via 'Heading 1' (no direct font) resolves to an
    effective font from the style/theme chain, not empty."""
    ev = _doc(lambda d: [d.add_heading("A", 1), d.add_heading("B", 1)])
    samples = ev["text_format_samples"].get("heading", [])
    assert samples and all(s["font"] for s in samples), "heading font should resolve via chain"
    assert all(s["size_pt"] for s in samples), "heading size should resolve via chain"


def test_direct_run_font_wins_over_style():
    from docx.shared import Pt

    def build(d):
        p = d.add_paragraph()
        r = p.add_run("explicit")
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)

    ev = _doc(build)
    body = ev["text_format_samples"].get("body", [])
    assert any(s["font"] == "Times New Roman" and s["size_pt"] == 13 for s in body)


# ---- declared / learned / undefined precedence ----

def test_undefined_text_format_is_called_out():
    """With no declared and no learnable text-format rule, the inspector emits a
    needs_review call-out rather than defaulting."""
    ev = {"text_format_samples": {"heading": [{"font": "Cambria", "size_pt": 14,
                                               "weight": "bold", "casing": "title"}]}}
    disc = BuildDiscipline()  # nothing declared, nothing learned yet
    findings = inspect_discipline(ev, disc)
    undef = [f for f in findings if f.key == "text_format.heading" and f.status == "needs_review"]
    assert undef, "undefined text-format rule must be called out"


def test_learned_rule_makes_consistent_doc_clean():
    ev = _doc(lambda d: [d.add_heading(f"S{i}", 1) for i in range(4)])
    disc = learn_discipline(BuildDiscipline(), ev)
    findings = inspect_discipline(ev, disc)
    drift = [f for f in findings if f.key.startswith("text_format.heading[") and f.status == "corrected"]
    assert not drift, "a consistent document should have no heading drift"
    assert disc.text_format["heading"].source == "learned"


def test_declared_rule_wins_over_learned():
    ev = _doc(lambda d: [d.add_heading("S", 1)])
    base = BuildDiscipline()
    base.text_format["heading"] = TextFormat(font="Arial", weight="bold", source="declared")
    learned = learn_discipline(base, ev)
    assert learned.text_format["heading"].font == "Arial"
    assert learned.text_format["heading"].source == "declared"


# ---- within-document drift ----

def test_drifting_heading_is_flagged():

    def build(d):
        for i in range(4):
            d.add_heading(f"S{i}", 1)
        p = d.add_paragraph()
        r = p.add_run("Drifted")
        r.font.name = "Times New Roman"
        r.bold = True
        p.style = d.styles["Heading 1"]

    ev = _doc(build)
    disc = learn_discipline(BuildDiscipline(), ev)
    findings = inspect_discipline(ev, disc)
    drift = [f for f in findings
             if f.key.startswith("text_format.heading[") and f.status == "corrected"
             and f.value["observed"] == "Times New Roman"]
    assert drift, "the drifting heading font should be flagged against the learned style"
