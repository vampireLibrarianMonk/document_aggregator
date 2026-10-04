"""OPT-IN DEVELOPER TOOL — not part of normal setup.

The app starts EMPTY by design: the only persistent content is what a user
creates (instantiating a sample case, generating, or uploading documents).
Nothing is seeded automatically, and no part of the runtime (API, worker,
Docker entrypoint, compose) invokes this script.

This exists only so a DEVELOPER can quickly populate a local AGGREGATION demo
(generated DOCX/PPTX/PDF/PNG ingested through the full pipeline, plus a few
supplementals) when manually testing the Ingestion / Search / Report tabs. It
mints a fresh proj_<id> each run, exactly like user-created content — it does
NOT reserve a special preloaded id.

    python backend/seed.py        # creates ONE throwaway aggregation project
"""
from __future__ import annotations

import io
import json
import sys
from pathlib import Path

# Allow "python backend/seed.py" from repo root.
sys.path.insert(0, str(Path(__file__).resolve().parent))

from app.aggregate import _sentiment  # noqa: E402
from app.models import (  # noqa: E402
    Project,
    Supplemental,  # noqa: E402
    SupplementalKind,
)
from app.pipeline import ingest_document  # noqa: E402
from app.store import store  # noqa: E402

REPO = Path(__file__).resolve().parents[1]
SAMPLES = REPO / "sample_docs"


def make_docx() -> bytes:
    import docx
    d = docx.Document()
    d.add_heading("TGX-9 Corrective Action Plan", level=0)
    d.add_heading("Background", level=1)
    d.add_paragraph(
        "Following repeated telemetry gateway outages in February and March 2026, "
        "a corrective action plan was established to address firmware 4.2.1 buffer "
        "management under thermal stress."
    )
    d.add_heading("Action Items", level=1)
    tbl = d.add_table(rows=1, cols=3)
    hdr = tbl.rows[0].cells
    hdr[0].text = "Action"
    hdr[1].text = "Owner"
    hdr[2].text = "Target Date"
    for action, owner, date in [
        ("Roll back affected sites to firmware 4.1.8", "Reliability Eng", "2026-03-20"),
        ("Validate 4.2.2 hotfix", "Firmware Team", "2026-04-05"),
        ("Add cabinet thermal alerting", "Site Ops", "2026-03-25"),
    ]:
        row = tbl.add_row().cells
        row[0].text, row[1].text, row[2].text = action, owner, date
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pptx() -> bytes:
    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "TGX-9 Incident Review"
    s.placeholders[1].text = "Executive briefing, March 2026"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Key Findings"
    tf = s2.placeholders[1].text_frame
    tf.text = "Firmware 4.2.1 buffer management fails under thermal stress"
    tf.add_paragraph().text = "Three sites affected; sites on 4.1.8 unaffected"
    tf.add_paragraph().text = "Manual power cycle used as workaround"
    if s2.has_notes_slide or True:
        s2.notes_slide.notes_text_frame.text = (
            "Presenter note: regional director escalated this after the last "
            "report shipped with incorrect tables."
        )
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_pdf() -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter)
    styles = getSampleStyleSheet()
    flow = [
        Paragraph("TGX-9 Executive Summary", styles["Title"]),
        Spacer(1, 12),
        Paragraph(
            "Between February and March 2026, three telemetry gateway sites running "
            "firmware 4.2.1 experienced correlated packet loss events under elevated "
            "cabinet temperatures. The North Ridge event on 2026-03-02 was the most "
            "severe, with 34 percent packet loss over four hours.",
            styles["BodyText"]),
        Spacer(1, 12),
        Paragraph("Recommended Path", styles["Heading1"]),
        Paragraph(
            "Roll back to firmware 4.1.8 or validate the 4.2.2 hotfix, and add "
            "cabinet thermal monitoring to the alerting baseline.",
            styles["BodyText"]),
    ]
    doc.build(flow)
    return buf.getvalue()


def make_png() -> bytes:
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (480, 200), "white")
    draw = ImageDraw.Draw(img)
    draw.rectangle([10, 10, 470, 190], outline="black", width=2)
    draw.text((30, 30), "Packet loss vs cabinet temp (scanned figure)", fill="black")
    draw.line([(40, 160), (120, 120), (200, 90), (300, 60), (440, 40)], fill="red", width=3)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def main() -> None:
    import uuid

    # Mint a fresh id like any user-created project (no reserved preload id).
    pid = "proj_" + uuid.uuid4().hex[:10]
    project = store.create_project(Project(id=pid, name="TGX-9 Incident Aggregation (dev seed)"))
    print(f"Created project {project.id} ({project.name})")

    # Static text samples from disk.
    corpus: list[tuple[str, bytes]] = []
    for name in ("field_report_2026-03-02.txt", "root_cause_notes_2026-03-15.md", "STYLE_GUIDE.md"):
        corpus.append((name, (SAMPLES / name).read_bytes()))

    # Generated binary samples (skip gracefully if a lib is missing).
    generators = {
        "corrective_action_plan.docx": make_docx,
        "incident_review.pptx": make_pptx,
        "executive_summary.pdf": make_pdf,
        "packet_loss_figure.png": make_png,
    }
    for fname, gen in generators.items():
        try:
            corpus.append((fname, gen()))
        except Exception as exc:  # missing optional lib
            print(f"  skip {fname}: {exc}")

    for fname, data in corpus:
        rec = ingest_document(project.id, fname, data)
        print(f"  ingested {fname:40s} -> {rec.overall_status:10s} "
              f"blocks={rec.block_count} chunks={rec.chunk_count} artifacts={rec.artifact_count}")

    # Supplementals (comments / angry emails / corrections / notes).
    supps = json.loads((SAMPLES / "supplementals.json").read_text(encoding="utf-8"))
    for s in supps:
        supp = Supplemental(
            id="supp_" + str(abs(hash(s["subject"])))[:10],
            kind=SupplementalKind(s["kind"]), author=s["author"],
            subject=s["subject"], body=s["body"],
            sentiment=_sentiment(s["subject"] + " " + s["body"]),
        )
        store.save_supplemental(project.id, supp)
        print(f"  supplemental [{supp.kind.value}/{supp.sentiment}] {supp.author}")

    print(f"\nDev seed complete: created aggregation project {project.id}. "
          "This is throwaway local data; the app does not seed on its own.")


if __name__ == "__main__":
    main()
