"""Author real DOCX/PPTX/PDF template + draft documents for every project.

Clients submit real documents, so we generate realistic ones (grounded in
common incident-report template structure: numbered sections, a management
briefing / corrective-actions table, sign-off fields, and a header/footer with
page numbers and a classification marking). Content comes from each project's
existing JSON draft/template so the converted result can be compared against the
JSON-authored baseline.

    python backend/build_sample_docs.py

Writes into sample_docs/project/<id>/first_attempt/generated/.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

SCEN = Path(__file__).resolve().parents[1] / "sample_docs" / "project"


def _load(sid: str, name: str) -> dict:
    return json.loads((SCEN / sid / "first_attempt" / name).read_text(encoding="utf-8"))


def _artifact_sections(artifact: dict, project: dict) -> list[dict]:
    """Normalize either a draft (has sections) or a template (has
    required_sections rubric) into a renderable section list."""
    if artifact.get("sections"):
        return artifact["sections"]
    # Template: render required sections as headings with blank field labels and
    # empty bodies/tables (a blank template a client would fill in).
    fields_by_section: dict[str, list[str]] = {}
    for f in project.get("fields", []):
        fields_by_section.setdefault(f["section"], []).append(f["key"])
    # Map required graphic id -> its managed name, so the template can show the
    # correctly-titled figure placeholder each section requires.
    gid_to_name = {g["graphic_id"]: g["name"] for g in _graphics_list(project)}
    out: list[dict] = []
    for spec in artifact.get("required_sections", []):
        sec: dict = {"key": spec["key"], "heading": spec["heading"],
                     "fields": {k: "" for k in fields_by_section.get(spec["key"], [])},
                     "body": "", "graphics": []}
        # A blank template still SHOWS where a figure belongs, correctly titled
        # and centered — this is the standard the draft is judged against.
        req_gid = spec.get("requires_graphic")
        if req_gid and req_gid in gid_to_name:
            sec["graphics"].append({"ref_name": gid_to_name[req_gid]})
        if spec.get("requires_table"):
            tspec = artifact.get("table_specs", {}).get(spec["requires_table"], {})
            sec["table"] = {"title": tspec.get("title", "").replace("{n}", "1"),
                            "font": tspec.get("font", ""),
                            "header_style": tspec.get("header_style", ""),
                            "columns": tspec.get("columns", []), "rows": []}
        out.append(sec)
    return out


def _graphics_list(project: dict) -> list[dict]:
    sid = str(project.get("id", ""))
    gjson = SCEN / sid / "corpus" / "graphics.json"
    if gjson.exists():
        return json.loads(gjson.read_text(encoding="utf-8")).get("graphics", [])
    return []


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

def build_docx(draft: dict, project: dict, path: Path) -> None:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    d = docx.Document()
    d.core_properties.title = draft.get("title", "")

    # Header + footer with page number field and a classification marking.
    section = d.sections[0]
    header = section.header
    header.paragraphs[0].text = draft.get("furniture", {}).get("header", {}).get("text", "") \
        or draft.get("title", "")
    footer_p = section.footer.paragraphs[0]
    classif = draft.get("furniture", {}).get("classification", "")
    footer_p.text = classif + "   " if classif else ""
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if draft.get("furniture", {}).get("page_numbers"):
        run = footer_p.add_run()
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        run._r.addprevious(fld)

    d.add_heading(draft.get("title", "Report"), level=0)

    for sec in _artifact_sections(draft, project):
        d.add_heading(sec["heading"], level=1)
        # fields as "Label: value" lines
        for fkey, fval in (sec.get("fields") or {}).items():
            label = fkey.replace("_", " ").title()
            d.add_paragraph(f"{label}: {fval}")
        if sec.get("body"):
            d.add_paragraph(sec["body"])
        for g in sec.get("graphics", []) or []:
            _embed_graphic(d, g, project)
        tbl = sec.get("table")
        if tbl and tbl.get("columns"):
            if tbl.get("title"):
                d.add_paragraph(tbl["title"])
            t = d.add_table(rows=1, cols=len(tbl["columns"]))
            for i, col in enumerate(tbl["columns"]):
                t.rows[0].cells[i].text = col
            for row in tbl.get("rows", []):
                cells = t.add_row().cells
                for i, val in enumerate(row):
                    if i < len(cells):
                        cells[i].text = str(val)

    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))


_GRAPHICS_CACHE: dict[str, dict] = {}


def _graphics_meta(project: dict) -> dict[str, dict]:
    """Load the project's managed figure metadata (name -> {file,title,w,h}).
    Cached per project id."""
    sid = str(project.get("id", ""))
    if sid in _GRAPHICS_CACHE:
        return _GRAPHICS_CACHE[sid]
    meta: dict[str, dict] = {}
    gjson = SCEN / sid / "corpus" / "graphics.json"
    if gjson.exists():
        data = json.loads(gjson.read_text(encoding="utf-8"))
        for g in data.get("graphics", []):
            meta[g["name"]] = g
    _GRAPHICS_CACHE[sid] = meta
    return meta


def _embed_graphic(d, g: dict, project: dict) -> None:
    """Embed a real PNG (centered, sized) with a centered title paragraph above
    it. Falls back to a text placeholder if the file is missing. The draft may
    override title/width to seed a defect (wrong/missing title, wrong size)."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches

    name = g.get("ref_name") or g.get("name") or ""
    meta = _graphics_meta(project).get(name, {})
    sid = str(project.get("id", ""))

    # Title: draft entry can override (or blank) it to seed a defect; otherwise
    # use the managed title from graphics.json.
    title = g.get("title", meta.get("title", ""))
    title_centered = g.get("title_centered", True)
    if title:
        tp = d.add_paragraph()
        tp.alignment = WD_ALIGN_PARAGRAPH.CENTER if title_centered else WD_ALIGN_PARAGRAPH.LEFT
        run = tp.add_run(title)
        run.bold = True

    # Image: embed the real file if present, else a named text placeholder so
    # the converter can still recover the reference.
    img_p = d.add_paragraph()
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER if g.get("align", "center") == "center" \
        else WD_ALIGN_PARAGRAPH.LEFT
    file_rel = meta.get("file")
    img_path = (SCEN / sid / "corpus" / file_rel) if file_rel else None
    if img_path and img_path.exists():
        # Width in inches: draft may shrink it to seed a size defect.
        width_px = g.get("width", meta.get("width", 480))
        run = img_p.add_run()
        run.add_picture(str(img_path), width=Inches(width_px / 96))
        # Carry the logical name as the drawing's docPr name so the converter
        # recovers the reference even though the media part is imageN.png.
        _set_docpr_name(img_p, name)
    else:
        run = img_p.add_run(f"[FIGURE: {name}]")
        run.italic = True


def _set_docpr_name(paragraph, name: str) -> None:
    """Set the drawing's docPr @name to the logical graphic name so the
    converter's _first_image_name recovers it."""
    from docx.oxml.ns import qn

    for docpr in paragraph._p.findall(".//" + qn("wp:docPr")):
        docpr.set("name", name)
        docpr.set("descr", name)


# --------------------------------------------------------------------------
# PPTX
# --------------------------------------------------------------------------

def build_pptx(draft: dict, project: dict, path: Path) -> None:
    from pptx import Presentation

    prs = Presentation()
    title_layout = prs.slide_layouts[0]
    body_layout = prs.slide_layouts[1]

    s = prs.slides.add_slide(title_layout)
    s.shapes.title.text = draft.get("title", "Report")
    classif = draft.get("furniture", {}).get("classification", "")
    s.placeholders[1].text = classif or ""

    for sec in _artifact_sections(draft, project):
        slide = prs.slides.add_slide(body_layout)
        slide.shapes.title.text = sec["heading"]
        tf = slide.placeholders[1].text_frame
        lines: list[str] = []
        for fkey, fval in (sec.get("fields") or {}).items():
            lines.append(f"{fkey.replace('_', ' ').title()}: {fval}")
        if sec.get("body"):
            lines.append(sec["body"])
        for g in sec.get("graphics", []) or []:
            lines.append(f"[FIGURE: {g['ref_name']}]")
        tf.text = lines[0] if lines else ""
        for extra in lines[1:]:
            tf.add_paragraph().text = extra
        tbl = sec.get("table")
        if tbl and tbl.get("columns"):
            note = slide.notes_slide.notes_text_frame
            note.text = "TABLE COLUMNS: " + " | ".join(tbl["columns"])

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def build_pdf(draft: dict, project: dict, path: Path) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table

    styles = getSampleStyleSheet()
    flow = [Paragraph(draft.get("title", "Report"), styles["Title"]), Spacer(1, 10)]
    for sec in _artifact_sections(draft, project):
        flow.append(Paragraph(sec["heading"], styles["Heading1"]))
        for fkey, fval in (sec.get("fields") or {}).items():
            flow.append(Paragraph(f"{fkey.replace('_', ' ').title()}: {fval}", styles["BodyText"]))
        if sec.get("body"):
            flow.append(Paragraph(sec["body"], styles["BodyText"]))
        for g in sec.get("graphics", []) or []:
            flow.append(Paragraph(f"[FIGURE: {g['ref_name']}]", styles["Italic"]))
        tbl = sec.get("table")
        if tbl and tbl.get("columns"):
            if tbl.get("title"):
                flow.append(Paragraph(tbl["title"], styles["Heading3"]))
            data = [tbl["columns"]] + [[str(c) for c in r] for r in tbl.get("rows", [])]
            flow.append(Table(data))
        flow.append(Spacer(1, 8))
    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(path), pagesize=letter).build(flow)


def build_docx_multipage(draft: dict, project: dict, pages: int, path: Path) -> None:
    """Render a multi-page DOCX: a page break between each page's section block,
    repeated header/footer with page numbers, denser content — closer to a real
    human submission. `draft` is a build_multipage draft (page-suffixed keys)."""
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    d = docx.Document()
    d.core_properties.title = draft.get("title", "")
    section = d.sections[0]
    section.header.paragraphs[0].text = draft.get("title", "Report")
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run("Page ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addprevious(fld)

    d.add_heading(draft.get("title", "Report"), level=0)

    current_page = 1
    for sec in draft.get("sections", []):
        key = sec["key"]
        page = int(key.split("__p")[1]) if "__p" in key else 1
        if page != current_page:
            d.add_page_break()
            current_page = page
        d.add_heading(sec.get("heading", key), level=1)
        for fkey, fval in (sec.get("fields") or {}).items():
            d.add_paragraph(f"{fkey.replace('_', ' ').title()}: {fval}")
        if sec.get("body"):
            d.add_paragraph(sec["body"])
        for g in sec.get("graphics", []) or []:
            p = d.add_paragraph()
            p.add_run(f"[FIGURE: {g['ref_name']}]").italic = True
        tbl = sec.get("table")
        if tbl and tbl.get("columns"):
            if tbl.get("title"):
                d.add_paragraph(tbl["title"])
            t = d.add_table(rows=1, cols=len(tbl["columns"]))
            for i, col in enumerate(tbl["columns"]):
                t.rows[0].cells[i].text = col
            for row in tbl.get("rows", []):
                cells = t.add_row().cells
                for i, val in enumerate(row):
                    if i < len(cells):
                        cells[i].text = str(val)

    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))


def build_pdf_multipage(draft: dict, project: dict, pages: int, path: Path) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table

    styles = getSampleStyleSheet()
    flow = [Paragraph(draft.get("title", "Report"), styles["Title"]), Spacer(1, 10)]
    current_page = 1
    for sec in draft.get("sections", []):
        key = sec["key"]
        page = int(key.split("__p")[1]) if "__p" in key else 1
        if page != current_page:
            flow.append(PageBreak())
            current_page = page
        flow.append(Paragraph(sec.get("heading", key), styles["Heading1"]))
        for fkey, fval in (sec.get("fields") or {}).items():
            flow.append(Paragraph(f"{fkey.replace('_', ' ').title()}: {fval}", styles["BodyText"]))
        if sec.get("body"):
            flow.append(Paragraph(sec["body"], styles["BodyText"]))
        for g in sec.get("graphics", []) or []:
            flow.append(Paragraph(f"[FIGURE: {g['ref_name']}]", styles["Italic"]))
        tbl = sec.get("table")
        if tbl and tbl.get("columns"):
            if tbl.get("title"):
                flow.append(Paragraph(tbl["title"], styles["Heading3"]))
            data = [tbl["columns"]] + [[str(c) for c in r] for r in tbl.get("rows", [])]
            flow.append(Table(data))
        flow.append(Spacer(1, 8))
    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(path), pagesize=letter).build(flow)


def _find_soffice() -> str | None:
    """Locate LibreOffice headless (Linux target). None if unavailable."""
    import shutil

    for name in ("soffice", "libreoffice"):
        path = shutil.which(name)
        if path:
            return path
    return None


def docx_to_pdf(docx_path: Path) -> Path | None:
    """Render a DOCX to PDF via LibreOffice headless, if present. Returns the
    PDF path, or None when soffice is unavailable (generation is skipped, never
    a hard dependency). Note: rendering to PDF discards DOCX style structure, so
    PDF inspection fidelity stays lower than DOCX regardless."""
    import subprocess

    soffice = _find_soffice()
    if not soffice:
        return None
    out_dir = docx_path.parent
    try:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir",
             str(out_dir), str(docx_path)],
            check=True, capture_output=True, timeout=120,
        )
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired):
        return None
    pdf = docx_path.with_suffix(".pdf")
    return pdf if pdf.exists() else None


def generate_multipage_samples(pages: int = 3, seed: int = 7) -> None:
    """Generate a multi-page DOCX + PDF for project 1 grown to `pages` pages."""
    import sys as _sys

    _sys.path.insert(0, str(Path(__file__).resolve().parent))
    from app import project as sc
    from app.knowledge.pagegrow import build_multipage

    b = build_multipage(sc.load_manifest("1"), sc.load_template("1"),
                        sc.load_first_attempt("1", "draft"), sc.load_corpus("1"),
                        sc.load_graphics("1"), sc.load_corrections("1"), pages, seed)
    out = SCEN / "1" / "first_attempt" / "generated" / "multipage"
    build_docx_multipage(b["draft"], b["project"], pages, out / f"draft_{pages}p.docx")
    build_pdf_multipage(b["draft"], b["project"], pages, out / f"draft_{pages}p.pdf")
    print(f"wrote {pages}-page draft.docx/.pdf ({len(b['draft']['sections'])} sections)")


def main() -> None:
    project_ids = [p.name for p in sorted(SCEN.iterdir()) if (p / "project.json").exists()]
    for sid in project_ids:
        project = json.loads((SCEN / sid / "project.json").read_text(encoding="utf-8"))
        out = SCEN / sid / "first_attempt" / "generated"
        for kind, jname in (("draft", "incident_report_draft.json"),
                            ("template", "incident_report_template.json")):
            src = _load(sid, jname)
            build_docx(src, project, out / f"{kind}.docx")
            build_pptx(src, project, out / f"{kind}.pptx")
            build_pdf(src, project, out / f"{kind}.pdf")
            print(f"project {sid}: wrote {kind}.docx/.pptx/.pdf")

    # Multi-page human-like samples for project 1 at a few page counts.
    for pages in (2, 3, 5):
        generate_multipage_samples(pages=pages, seed=7)


if __name__ == "__main__":
    main()
