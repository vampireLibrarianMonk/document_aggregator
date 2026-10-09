"""Parser registry / router.

Each parser converts source bytes into a list of canonical Blocks plus a list
of Artifacts. Optional libraries (python-docx, python-pptx, pypdf, Pillow) are
imported lazily so the pipeline degrades gracefully: if a parser's dependency
is missing, it reports that clearly instead of crashing the service.

File-type routing is by MIME/extension; a real build would sniff content.
"""
from __future__ import annotations

import io
import re
import uuid
from dataclasses import dataclass, field

from .models import Artifact, Block, Provenance
from .ocr import ocr_available, ocr_image_bytes, ocr_pdf_bytes


@dataclass
class ParseResult:
    blocks: list[Block] = field(default_factory=list)
    artifacts: list[Artifact] = field(default_factory=list)
    parser: str = "unknown"
    parser_version: str = ""
    method: str = "native"
    notes: str = ""
    # Extracted embedded image bytes, keyed by a logical filename (docx/pptx).
    # Each Artifact(type=image) carries metadata["image_name"] into this map so
    # a draft/template's own figures can be served + anchored to their text.
    images: dict[str, bytes] = field(default_factory=dict)


def _bid() -> str:
    return "block_" + uuid.uuid4().hex[:12]


def _aid() -> str:
    return "artifact_" + uuid.uuid4().hex[:12]


def _prov(parser: str, version: str, method: str, doc_id: str) -> Provenance:
    return Provenance(parser=parser, parser_version=version, method=method, source_document_id=doc_id)


# --------------------------------------------------------------------------
# TXT / Markdown
# --------------------------------------------------------------------------

def parse_text(data: bytes, doc_id: str, is_markdown: bool) -> ParseResult:
    text = data.decode("utf-8", errors="replace")
    parser = "markdown_native" if is_markdown else "text_native"
    prov = _prov(parser, "1", "native", doc_id)
    blocks: list[Block] = []
    order = 0
    section_path: list[str] = []
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        btype = "paragraph"
        if is_markdown and line.lstrip().startswith("#"):
            btype = "heading"
            section_path = [line.lstrip("# ").strip()]
        elif is_markdown and line.lstrip()[:2] in ("- ", "* "):
            btype = "list_item"
        blocks.append(Block(
            id=_bid(), type=btype, text=line.strip(),
            reading_order=order, section_path=list(section_path), provenance=prov,
        ))
        order += 1
    return ParseResult(blocks=blocks, parser=parser, parser_version="1", notes="")


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

def parse_docx(data: bytes, doc_id: str) -> ParseResult:
    try:
        import docx  # python-docx
    except ImportError:
        return ParseResult(parser="docx", notes="python-docx not installed; DOCX parsing skipped")

    document = docx.Document(io.BytesIO(data))
    prov = _prov("python-docx", getattr(docx, "__version__", "?"), "native", doc_id)
    blocks: list[Block] = []
    artifacts: list[Artifact] = []
    images: dict[str, bytes] = {}
    order = 0
    section_path: list[str] = []

    # Map relationship-id -> image bytes for the main document part, so an inline
    # image's r:embed can be resolved to its actual bytes.
    rels = getattr(document.part, "rels", {})

    def _image_from_rid(rid: str):
        try:
            part = rels[rid].target_part
            return getattr(part, "blob", None), getattr(part, "partname", "")
        except Exception:
            return None, ""

    def _para_image_rids(para) -> list[tuple[str, str, str]]:
        """Return (rid, docpr_name, docpr_descr) for each inline/anchored image
        in a paragraph, in document order."""
        from docx.oxml.ns import qn
        out: list[tuple[str, str, str]] = []
        p = para._p
        for drawing in p.iter(qn("w:drawing")):
            # docPr carries a human name/description (often the figure name).
            name = descr = ""
            for docpr in drawing.iter(qn("wp:docPr")):
                name = docpr.get("name", "") or ""
                descr = docpr.get("descr", "") or ""
                break
            for blip in drawing.iter(qn("a:blip")):
                rid = blip.get(qn("r:embed")) or blip.get(qn("r:link"))
                if rid:
                    out.append((rid, name, descr))
        return out

    img_seq = 0
    last_block_id: str | None = None
    for para in document.paragraphs:
        txt = para.text.strip()
        img_rids = _para_image_rids(para)

        # Emit a text block for the paragraph (if it has text). A paragraph may
        # hold BOTH text (a caption) and an image; keep the text as a block and
        # anchor the image to it.
        this_block_id: str | None = None
        if txt:
            style = (para.style.name or "").lower() if para.style else ""
            if "heading" in style or "title" in style:
                btype = "heading"
                section_path = [txt]
            elif "list" in style:
                btype = "list_item"
            elif img_rids:
                btype = "caption"  # text riding with an image reads as a caption
            else:
                btype = "paragraph"
            bid = _bid()
            blocks.append(Block(id=bid, type=btype, text=txt, reading_order=order,
                                section_path=list(section_path), provenance=prov))
            this_block_id = bid
            last_block_id = bid
            order += 1

        # Emit an image artifact per embedded image, anchored to the paragraph it
        # sits in (or the preceding text block if the image paragraph is empty).
        for rid, docpr_name, docpr_descr in img_rids:
            blob, partname = _image_from_rid(rid)
            img_seq += 1
            ext = str(partname).rsplit(".", 1)[-1].lower() if "." in str(partname) else "png"
            if ext not in ("png", "jpg", "jpeg", "gif", "bmp", "tiff"):
                ext = "png"
            # A stable logical name: prefer the docPr name, else figure_N.
            base = (docpr_name or docpr_descr or f"figure_{img_seq}").strip()
            base = re.sub(r"[^A-Za-z0-9._-]+", "_", base).strip("_") or f"figure_{img_seq}"
            if not base.lower().endswith((".png", ".jpg", ".jpeg", ".gif")):
                base = f"{base}.{ext}"
            if blob:
                images[base] = blob
            anchor = this_block_id or last_block_id
            artifacts.append(Artifact(
                id=_aid(), type="image", classification="inline_image",
                description=(docpr_descr or docpr_name or None),
                anchor_block_id=anchor,
                nearby_block_ids=[b for b in [last_block_id] if b and b != anchor],
                metadata={"image_name": base, "section_path": list(section_path),
                          "reading_order": order},
            ))

    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        flat = " | ".join(" ".join(r) for r in rows)
        blocks.append(Block(id=_bid(), type="table", text=flat, reading_order=order,
                            section_path=list(section_path), provenance=prov, table=rows))
        order += 1

    return ParseResult(blocks=blocks, artifacts=artifacts, images=images,
                       parser="python-docx", parser_version=getattr(docx, "__version__", "?"))


# --------------------------------------------------------------------------
# PPTX
# --------------------------------------------------------------------------

def parse_pptx(data: bytes, doc_id: str) -> ParseResult:
    try:
        from pptx import Presentation
    except ImportError:
        return ParseResult(parser="pptx", notes="python-pptx not installed; PPTX parsing skipped")

    prs = Presentation(io.BytesIO(data))
    prov = _prov("python-pptx", "1", "native", doc_id)
    blocks: list[Block] = []
    artifacts: list[Artifact] = []
    images: dict[str, bytes] = {}
    order = 0
    img_seq = 0
    for sidx, slide in enumerate(prs.slides, start=1):
        section_path = [f"Slide {sidx}"]
        slide_last_block: str | None = None
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                txt = shape.text_frame.text.strip()
                is_title = bool(getattr(shape, "is_placeholder", False)) and order == 0
                bid = _bid()
                blocks.append(Block(id=bid, type="slide_title" if is_title else "paragraph",
                                    text=txt, slide=sidx, reading_order=order,
                                    section_path=list(section_path), provenance=prov))
                slide_last_block = bid
                order += 1
            if shape.shape_type == 13:  # PICTURE
                img_seq += 1
                name = f"slide{sidx}_figure_{img_seq}.png"
                try:
                    img = shape.image
                    blob = img.blob
                    ext = (img.ext or "png").lower()
                    name = f"slide{sidx}_figure_{img_seq}.{ext}"
                    images[name] = blob
                except Exception:
                    pass
                # Anchor the image to the most recent text block on this slide
                # (its title/caption), so it travels with the right text.
                artifacts.append(Artifact(
                    id=_aid(), type="image", slide=sidx, classification="slide_image",
                    anchor_block_id=slide_last_block,
                    metadata={"image_name": name, "section_path": list(section_path),
                              "reading_order": order}))
        # speaker notes
        if slide.has_notes_slide:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                blocks.append(Block(id=_bid(), type="notes", text=note, slide=sidx,
                                    reading_order=order, section_path=list(section_path),
                                    provenance=prov))
                order += 1
    return ParseResult(blocks=blocks, artifacts=artifacts, images=images,
                       parser="python-pptx", parser_version="1")


# --------------------------------------------------------------------------
# PDF (born-digital)
# --------------------------------------------------------------------------

def parse_pdf(data: bytes, doc_id: str) -> ParseResult:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ParseResult(parser="pdf", notes="pypdf not installed; PDF parsing skipped")

    reader = PdfReader(io.BytesIO(data))
    prov = _prov("pypdf", "1", "native", doc_id)
    blocks: list[Block] = []
    order = 0
    total_text = 0
    for pidx, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        total_text += len(text)
        for para in [p.strip() for p in text.split("\n") if p.strip()]:
            blocks.append(Block(id=_bid(), type="paragraph", text=para, page=pidx,
                                reading_order=order, provenance=prov))
            order += 1
    if total_text >= 20:
        return ParseResult(blocks=blocks, parser="pypdf", parser_version="1",
                           method="native_pdf", notes="")

    # Little/no native text: scanned/flattened PDF. Try OCR when enabled + an
    # engine is installed; otherwise report needs_ocr exactly as before.
    if ocr_available():
        pages = ocr_pdf_bytes(data)
        ocr_blocks: list[Block] = []
        o = 0
        for pidx, page_text in enumerate(pages, start=1):
            for para in [p.strip() for p in (page_text or "").split("\n") if p.strip()]:
                ocr_blocks.append(Block(id=_bid(), type="paragraph", text=para, page=pidx,
                                        reading_order=o,
                                        provenance=_prov("pypdf+ocr", "1", "ocr_pdf", doc_id)))
                o += 1
        if ocr_blocks:
            return ParseResult(blocks=ocr_blocks, parser="pypdf+ocr", parser_version="1",
                               method="ocr_pdf",
                               notes=f"Scanned PDF recovered via OCR ({len(ocr_blocks)} text blocks).")
        return ParseResult(blocks=blocks, parser="pypdf", parser_version="1",
                           method="needs_ocr",
                           notes="Scanned PDF; OCR produced no text (empty or unreadable pages).")
    return ParseResult(blocks=blocks, parser="pypdf", parser_version="1", method="needs_ocr",
                       notes="Little/no native text extracted; scanned PDF -> OCR escalation "
                             "required (OCR disabled or no engine installed).")


# --------------------------------------------------------------------------
# Images (PNG/JPEG/TIFF) -> artifact + EXIF timestamp evidence
# --------------------------------------------------------------------------

def parse_image(data: bytes, doc_id: str, mime: str) -> ParseResult:
    art = Artifact(id=_aid(), type="image", classification="standalone_image",
                   content_hash=None, metadata={"mime": mime})
    notes = "Image ingested as artifact. Baseline keeps location/relationship; OCR/VLM optional."
    try:
        from PIL import Image
        from PIL.ExifTags import TAGS
        img = Image.open(io.BytesIO(data))
        art.metadata["width"], art.metadata["height"] = img.size
        exif = getattr(img, "_getexif", lambda: None)() or {}
        for tag_id, val in exif.items():
            name = TAGS.get(tag_id, str(tag_id))
            if name in ("DateTimeOriginal", "DateTime"):
                art.metadata["exif_datetime_original"] = str(val)
    except ImportError:
        notes += " (Pillow not installed; no dimensions/EXIF)"
    except Exception as exc:  # malformed image should not crash the worker
        notes += f" (image inspect failed: {exc})"

    # OCR the image to text when enabled + an engine is installed, so a photo /
    # scan of a document yields usable content (not just an artifact). Keeps the
    # artifact + EXIF either way.
    blocks: list[Block] = []
    method = "native"
    if ocr_available():
        text = ocr_image_bytes(data)
        if text:
            prov = _prov("image+ocr", "1", "ocr_image", doc_id)
            for order, para in enumerate(p.strip() for p in text.split("\n") if p.strip()):
                blocks.append(Block(id=_bid(), type="paragraph", text=para,
                                    reading_order=order, provenance=prov))
            method = "ocr_image"
            notes += f" OCR recovered {len(blocks)} text block(s)."
        else:
            notes += " OCR enabled but produced no text."
    return ParseResult(blocks=blocks, artifacts=[art], parser="image_native",
                       parser_version="1", method=method, notes=notes)


# --------------------------------------------------------------------------
# Router
# --------------------------------------------------------------------------

def route_and_parse(filename: str, mime: str, data: bytes, doc_id: str) -> ParseResult:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext in ("txt",) or mime == "text/plain":
        return parse_text(data, doc_id, is_markdown=False)
    if ext in ("md", "markdown"):
        return parse_text(data, doc_id, is_markdown=True)
    if ext == "docx" or "wordprocessingml" in mime:
        return parse_docx(data, doc_id)
    if ext == "pptx" or "presentationml" in mime:
        return parse_pptx(data, doc_id)
    if ext == "pdf" or mime == "application/pdf":
        return parse_pdf(data, doc_id)
    if ext in ("png", "jpg", "jpeg", "tif", "tiff") or mime.startswith("image/"):
        return parse_image(data, doc_id, mime)
    # Fallback: try decoding as text (Tika would sit here in a full build).
    return parse_text(data, doc_id, is_markdown=False)
