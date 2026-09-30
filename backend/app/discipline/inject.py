"""Discipline violation injection for the inspection alpha loop.

Builds a DOCX that is discipline-compliant by default, then applies a named
placement/format violation so tests can confirm the inspector detects exactly
that violation. This exercises the app's core purpose (document inspection) with
ground truth about what was broken.
"""
from __future__ import annotations

import io

VIOLATIONS = (
    "missing_caption",
    "caption_above",
    "missing_table_title",
    "table_title_below",
    "missing_page_numbers",
    "missing_header",
    "wrong_table_header_style",
)


def build_docx_with_violation(violation: str | None) -> bytes:
    """Return DOCX bytes: a compliant incident page, with one violation applied
    (or fully compliant when violation is None)."""
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    d = docx.Document()
    d.core_properties.title = "Inspection Test Report"

    # Header (unless suppressed).
    if violation != "missing_header":
        d.sections[0].header.paragraphs[0].text = "Inspection Test Report"

    # Footer with page number (unless suppressed).
    if violation != "missing_page_numbers":
        fp = d.sections[0].footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run("Page ")
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        run._r.addprevious(fld)

    d.add_heading("1. Description", level=1)
    d.add_paragraph("The system experienced a fault during the window.")

    # Figure with caption. Caption position controlled by violation.
    def add_figure():
        d.add_paragraph("[FIGURE: system_diagram.png]")

    def add_caption():
        d.add_paragraph("Figure 1: System diagram")

    if violation == "missing_caption":
        add_figure()
    elif violation == "caption_above":
        add_caption()
        add_figure()
    else:  # compliant: caption below
        add_figure()
        add_caption()

    d.add_heading("2. Corrective Actions", level=1)

    # Table with a title. Title presence/position controlled by violation.
    def add_table_title():
        d.add_paragraph("Table 1: Corrective Actions")

    def add_table(header_style_bold: bool):
        t = d.add_table(rows=1, cols=3)
        hdr = t.rows[0].cells
        for i, col in enumerate(["Action", "Owner", "Due Date"]):
            hdr[i].text = col
            if header_style_bold:
                for para in hdr[i].paragraphs:
                    for run in para.runs:
                        run.font.bold = True
                    if not para.runs:
                        para.add_run(col).font.bold = True
        row = t.add_row().cells
        row[0].text, row[1].text, row[2].text = "Do the thing", "Owner A", "2026-01-01"

    if violation == "missing_table_title":
        add_table(header_style_bold=True)
    elif violation == "table_title_below":
        add_table(header_style_bold=True)
        add_table_title()
    elif violation == "wrong_table_header_style":
        add_table_title()
        add_table(header_style_bold=False)  # not the table-header style
    else:  # compliant: title above, bold header
        add_table_title()
        add_table(header_style_bold=True)

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
