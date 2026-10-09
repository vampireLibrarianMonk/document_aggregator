"""Exporters: render the intermediate report JSON into deliverable formats.

JSON and Markdown are always available (stdlib). DOCX/PPTX/PDF require optional
libraries and raise a clear error if the library is missing, rather than
failing opaquely. Rendering is a formatting step over already-assembled
content — nothing is generated here.
"""
from __future__ import annotations

import io
import json


def export_json(report: dict) -> tuple[bytes, str, str]:
    data = json.dumps(report, indent=2, ensure_ascii=False).encode("utf-8")
    return data, "application/json", "report.json"


def export_markdown(report: dict) -> tuple[bytes, str, str]:
    lines: list[str] = []
    proj = report.get("project", {})
    lines.append(f"# Aggregated Report — {proj.get('name', '')}")
    lines.append("")
    s = report.get("summary", {})
    lines.append(f"- Source documents: {s.get('source_documents', 0)}")
    lines.append(f"- Supplementals: {s.get('supplementals_total', 0)} "
                 f"({s.get('negative_supplementals', 0)} negative)")
    lines.append(f"- Ordering: {report.get('ordering', '')}")
    lines.append("")
    for sec in report.get("sections", []):
        lines.append(f"## {sec['title']}")
        lines.append(f"*Effective DTG: {sec.get('effective_dtg')} "
                     f"(source: {sec.get('effective_dtg_source')})*")
        lines.append("")
        for item in sec.get("content", []):
            if item.get("table"):
                for row in item["table"]:
                    lines.append("| " + " | ".join(row) + " |")
                lines.append("")
            elif item["type"] == "heading":
                lines.append(f"### {item['text']}")
            elif item["type"] == "list_item":
                lines.append(f"- {item['text']}")
            else:
                lines.append(item["text"])
            lines.append("")
        if sec.get("supplementals"):
            lines.append("**Supplementals:**")
            for sup in sec["supplementals"]:
                lines.append(f"- ({sup['kind']}/{sup['sentiment']}) {sup['author']}: {sup['body']}")
            lines.append("")
    return "\n".join(lines).encode("utf-8"), "text/markdown", "report.md"


def export_docx(report: dict) -> tuple[bytes, str, str]:
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError("python-docx not installed; cannot export DOCX") from exc
    document = docx.Document()
    proj = report.get("project", {})
    document.add_heading(f"Aggregated Report — {proj.get('name', '')}", level=0)
    for sec in report.get("sections", []):
        document.add_heading(sec["title"], level=1)
        dtg_para = document.add_paragraph(
            f"Effective DTG: {sec.get('effective_dtg')} (source: {sec.get('effective_dtg_source')})"
        )
        # Italic is a run-level property in python-docx, not a paragraph one.
        if dtg_para.runs:
            dtg_para.runs[0].italic = True
        for item in sec.get("content", []):
            if item.get("table") and item["table"]:
                rows = item["table"]
                tbl = document.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row):
                        tbl.cell(ri, ci).text = cell
            elif item["type"] == "heading":
                document.add_heading(item["text"], level=2)
            elif item["type"] == "list_item":
                document.add_paragraph(item["text"], style="List Bullet")
            else:
                document.add_paragraph(item["text"])
        for sup in sec.get("supplementals", []):
            document.add_paragraph(
                f"[{sup['kind']}/{sup['sentiment']}] {sup['author']}: {sup['body']}"
            )
    buf = io.BytesIO()
    document.save(buf)
    return (buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "report.docx")


def export_pptx(report: dict) -> tuple[bytes, str, str]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("python-pptx not installed; cannot export PPTX") from exc
    prs = Presentation()
    title_layout = prs.slide_layouts[0]
    body_layout = prs.slide_layouts[1]
    proj = report.get("project", {})
    s = prs.slides.add_slide(title_layout)
    s.shapes.title.text = f"Aggregated Report — {proj.get('name', '')}"
    s.placeholders[1].text = report.get("ordering", "")
    for sec in report.get("sections", []):
        slide = prs.slides.add_slide(body_layout)
        slide.shapes.title.text = sec["title"]
        body = slide.placeholders[1].text_frame
        body.text = f"Effective DTG: {sec.get('effective_dtg')}"
        for item in sec.get("content", [])[:12]:
            p = body.add_paragraph()
            p.text = (item.get("text", "") or "")[:180]
    buf = io.BytesIO()
    prs.save(buf)
    return (buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "report.pptx")


def export_pdf(report: dict) -> tuple[bytes, str, str]:
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table
    except ImportError as exc:
        raise RuntimeError("reportlab not installed; cannot export PDF") from exc
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter)
    styles = getSampleStyleSheet()
    flow = []
    proj = report.get("project", {})
    flow.append(Paragraph(f"Aggregated Report — {proj.get('name', '')}", styles["Title"]))
    flow.append(Spacer(1, 12))
    for sec in report.get("sections", []):
        flow.append(Paragraph(sec["title"], styles["Heading1"]))
        flow.append(Paragraph(
            f"Effective DTG: {sec.get('effective_dtg')} "
            f"(source: {sec.get('effective_dtg_source')})", styles["Italic"]))
        for item in sec.get("content", []):
            if item.get("table") and item["table"]:
                flow.append(Table(item["table"]))
            else:
                style = styles["Heading2"] if item["type"] == "heading" else styles["BodyText"]
                text = (item.get("text", "") or "").replace("<", "&lt;").replace(">", "&gt;")
                if text:
                    flow.append(Paragraph(text, style))
        for sup in sec.get("supplementals", []):
            flow.append(Paragraph(
                f"[{sup['kind']}/{sup['sentiment']}] {sup['author']}: {sup['body']}",
                styles["BodyText"]))
        flow.append(Spacer(1, 12))
    doc.build(flow)
    return buf.getvalue(), "application/pdf", "report.pdf"


EXPORTERS = {
    "json": export_json,
    "markdown": export_markdown,
    "docx": export_docx,
    "pptx": export_pptx,
    "pdf": export_pdf,
}


def export_report(report: dict, fmt: str) -> tuple[bytes, str, str]:
    fn = EXPORTERS.get(fmt)
    if not fn:
        raise ValueError(f"unsupported export format: {fmt}")
    return fn(report)


# --------------------------------------------------------------------------
# Corrected-report exporters
#
# These render the CORRECTED intermediate report (the Correction Pipeline's
# output from reconcile()), not the aggregated corpus dump. The deliverable is
# the finished, source-grounded report: each field shows its final value, a
# conflict is shown as unresolved with its candidates, the table is filled, and
# the figures are placed. Manual resolutions are already reflected because the
# report is produced by run_reconciliation over the merged corrections.
# --------------------------------------------------------------------------

def _cr_field_line(f: dict) -> str:
    """One field rendered in document voice: 'Label: value' / unresolved / blank
    with a short status note when it is not an untouched value."""
    label = f.get("label", f.get("key", ""))
    status = f.get("status", "")
    if status == "conflict":
        cands = f.get("candidates", [])
        choices = " vs ".join(str(c.get("value")) for c in cands) or "unresolved"
        return f"{label}: [unresolved conflict — choose: {choices}]"
    if status == "needs_review":
        return f"{label}: [needs review — no source value]"
    val = f.get("value")
    return f"{label}: {'' if val is None else val}"


def _cr_sections(report: dict) -> list[dict]:
    return report.get("sections", [])


def _cr_title(report: dict) -> str:
    return report.get("title", "Corrected Report")


def _cr_status_summary(report: dict) -> str:
    s = report.get("summary", {})
    parts = [f"{k.replace('_', ' ')}: {v}" for k, v in s.items() if k != "total_units"]
    total = s.get("total_units")
    head = f"{total} units" if total is not None else ""
    return "  ·  ".join([p for p in [head, *parts] if p])


def _figure_path(figures_dir, name: str):
    """Resolve a figure file inside figures_dir if it exists, else None."""
    if not figures_dir or not name:
        return None
    from pathlib import Path
    p = Path(figures_dir) / name
    return p if p.exists() and p.is_file() else None


def _report_figures(report: dict) -> list[dict]:
    """All graphics across the report's sections, in reading order."""
    out: list[dict] = []
    for sec in _cr_sections(report):
        out.extend(sec.get("graphics", []))
    return out


def export_corrected_json(report: dict, figures_dir=None) -> tuple[bytes, str, str]:
    data = json.dumps(report, indent=2, ensure_ascii=False).encode("utf-8")
    return data, "application/json", "corrected_report.json"


def _corrected_markdown_text(report: dict, embed_figures: bool) -> str:
    lines: list[str] = [f"# {_cr_title(report)}", ""]
    summ = _cr_status_summary(report)
    if summ:
        lines += [f"*{summ}*", ""]
    for sec in _cr_sections(report):
        lines.append(f"## {sec.get('heading', sec.get('key', ''))}")
        lines.append("")
        for f in sec.get("fields", []):
            lines.append(f"- {_cr_field_line(f)}")
        for t in sec.get("tables", []):
            title = t.get("title", "Table").replace("{n}", str(t.get("table_number") or 1))
            lines += ["", f"**{title}**", ""]
            cols = t.get("columns", [])
            if cols:
                lines.append("| " + " | ".join(cols) + " |")
                lines.append("| " + " | ".join(["---"] * len(cols)) + " |")
                for row in t.get("rows", []):
                    cells = ["needs review" if c == "[needs_review]" else str(c) for c in row]
                    lines.append("| " + " | ".join(cells) + " |")
                if not t.get("rows"):
                    lines.append("| " + " | ".join(["(filled from corpus)"] * len(cols)) + " |")
        for g in sec.get("graphics", []):
            num = f"Figure {g['figure_number']}: " if g.get("figure_number") else ""
            caption = f"{num}{g.get('caption') or g.get('name', '')}"
            name = g.get("name", "")
            if embed_figures and name:
                # GitHub-style embedded image (relative to the markdown file).
                lines += ["", f"![{caption}](figures/{name})", "", f"*{caption}*"]
            else:
                lines += ["", f"*{caption}* (`{name}`)"]
        lines.append("")
    fe = report.get("furniture", {}).get("elements", [])
    if fe:
        lines += ["## Page elements", ""]
        for f in fe:
            lines.append(f"- {_cr_field_line(f)}")
        lines.append("")
    df = report.get("discipline_findings", [])
    if df:
        lines += ["## Formatting & placement checks", ""]
        for f in df:
            lines.append(f"- {_cr_field_line(f)}")
        lines.append("")
    return "\n".join(lines)


def export_corrected_markdown(report: dict, figures_dir=None) -> tuple[bytes, str, str]:
    """Markdown with GitHub-style embedded images, delivered as a ZIP archive
    containing corrected_report.md + a figures/ folder with the real PNGs, so
    the images render when the archive is unpacked (as they do in a GitHub
    README). If no figures are available, still returns a ZIP for consistency."""
    import zipfile

    figs = _report_figures(report)
    md = _corrected_markdown_text(report, embed_figures=True)

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("corrected_report.md", md)
        seen: set[str] = set()
        for g in figs:
            name = g.get("name", "")
            if not name or name in seen:
                continue
            seen.add(name)
            fp = _figure_path(figures_dir, name)
            if fp is not None:
                z.write(fp, f"figures/{name}")
    return buf.getvalue(), "application/zip", "corrected_report.zip"


def export_corrected_docx(report: dict, figures_dir=None) -> tuple[bytes, str, str]:
    try:
        import docx
        from docx.shared import Inches
    except ImportError as exc:
        raise RuntimeError("python-docx not installed; cannot export DOCX") from exc
    d = docx.Document()
    d.add_heading(_cr_title(report), level=0)
    summ = _cr_status_summary(report)
    if summ:
        p = d.add_paragraph(summ)
        if p.runs:
            p.runs[0].italic = True
    for sec in _cr_sections(report):
        d.add_heading(sec.get("heading", sec.get("key", "")), level=1)
        for f in sec.get("fields", []):
            d.add_paragraph(_cr_field_line(f))
        for t in sec.get("tables", []):
            title = t.get("title", "Table").replace("{n}", str(t.get("table_number") or 1))
            d.add_heading(title, level=2)
            cols = t.get("columns", [])
            rows = t.get("rows", [])
            if cols:
                tbl = d.add_table(rows=1 + len(rows), cols=len(cols))
                for ci, c in enumerate(cols):
                    tbl.cell(0, ci).text = c
                for ri, row in enumerate(rows, start=1):
                    for ci, cell in enumerate(row):
                        tbl.cell(ri, ci).text = (
                            "needs review" if cell == "[needs_review]" else str(cell))
        for g in sec.get("graphics", []):
            num = f"Figure {g['figure_number']}: " if g.get("figure_number") else ""
            caption = f"{num}{g.get('caption') or g.get('name', '')}"
            # Embed the REAL image when the file is available; else caption only.
            fp = _figure_path(figures_dir, g.get("name", ""))
            if fp is not None:
                try:
                    d.add_picture(str(fp), width=Inches(5.5))
                except Exception:
                    pass
            cap = d.add_paragraph(caption)
            if cap.runs:
                cap.runs[0].italic = True
    fe = report.get("furniture", {}).get("elements", [])
    if fe:
        d.add_heading("Page elements", level=1)
        for f in fe:
            d.add_paragraph(_cr_field_line(f))
    df = report.get("discipline_findings", [])
    if df:
        d.add_heading("Formatting & placement checks", level=1)
        for f in df:
            d.add_paragraph(_cr_field_line(f))
    buf = io.BytesIO()
    d.save(buf)
    return (buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "corrected_report.docx")


def export_corrected_pptx(report: dict, figures_dir=None) -> tuple[bytes, str, str]:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        raise RuntimeError("python-pptx not installed; cannot export PPTX") from exc
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = _cr_title(report)
    s.placeholders[1].text = _cr_status_summary(report)
    body_layout = prs.slide_layouts[1]
    for sec in _cr_sections(report):
        slide = prs.slides.add_slide(body_layout)
        slide.shapes.title.text = sec.get("heading", sec.get("key", ""))
        tf = slide.placeholders[1].text_frame
        lines = [_cr_field_line(f) for f in sec.get("fields", [])]
        for t in sec.get("tables", []):
            lines.append(t.get("title", "Table").replace("{n}", str(t.get("table_number") or 1)))
        tf.text = lines[0] if lines else ""
        for extra in lines[1:]:
            tf.add_paragraph().text = extra
        # Embed each section figure on its own slide so the real image shows.
        for g in sec.get("graphics", []):
            caption = g.get("caption") or g.get("name", "")
            fp = _figure_path(figures_dir, g.get("name", ""))
            fig_slide = prs.slides.add_slide(prs.slide_layouts[5])
            num = f"Figure {g['figure_number']}: " if g.get("figure_number") else ""
            fig_slide.shapes.title.text = f"{num}{caption}"
            if fp is not None:
                try:
                    fig_slide.shapes.add_picture(str(fp), Inches(1), Inches(1.8), width=Inches(8))
                except Exception:
                    pass
    buf = io.BytesIO()
    prs.save(buf)
    return (buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "corrected_report.pptx")


def export_corrected_pdf(report: dict, figures_dir=None) -> tuple[bytes, str, str]:
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table
    except ImportError as exc:
        raise RuntimeError("reportlab not installed; cannot export PDF") from exc
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter)
    styles = getSampleStyleSheet()

    def esc(s: str) -> str:
        return str(s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    flow = [Paragraph(esc(_cr_title(report)), styles["Title"])]
    summ = _cr_status_summary(report)
    if summ:
        flow.append(Paragraph(esc(summ), styles["Italic"]))
    flow.append(Spacer(1, 12))
    for sec in _cr_sections(report):
        flow.append(Paragraph(esc(sec.get("heading", sec.get("key", ""))), styles["Heading1"]))
        for f in sec.get("fields", []):
            flow.append(Paragraph(esc(_cr_field_line(f)), styles["BodyText"]))
        for t in sec.get("tables", []):
            title = t.get("title", "Table").replace("{n}", str(t.get("table_number") or 1))
            flow.append(Paragraph(esc(title), styles["Heading2"]))
            cols = t.get("columns", [])
            rows = t.get("rows", [])
            if cols:
                grid = [cols] + [
                    ["needs review" if c == "[needs_review]" else str(c) for c in row]
                    for row in rows
                ]
                flow.append(Table(grid))
        for g in sec.get("graphics", []):
            num = f"Figure {g['figure_number']}: " if g.get("figure_number") else ""
            caption = f"{num}{g.get('caption') or g.get('name', '')}"
            fp = _figure_path(figures_dir, g.get("name", ""))
            if fp is not None:
                try:
                    flow.append(Image(str(fp), width=400, height=225))
                except Exception:
                    pass
            flow.append(Paragraph(esc(caption), styles["Italic"]))
        flow.append(Spacer(1, 10))
    doc.build(flow)
    return buf.getvalue(), "application/pdf", "corrected_report.pdf"


CORRECTED_EXPORTERS = {
    "json": export_corrected_json,
    "markdown": export_corrected_markdown,
    "docx": export_corrected_docx,
    "pptx": export_corrected_pptx,
    "pdf": export_corrected_pdf,
}


def export_corrected_report(report: dict, fmt: str, figures_dir=None) -> tuple[bytes, str, str]:
    fn = CORRECTED_EXPORTERS.get(fmt)
    if not fn:
        raise ValueError(f"unsupported export format: {fmt}")
    return fn(report, figures_dir)
